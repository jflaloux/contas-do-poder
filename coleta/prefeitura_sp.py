"""Prefeitura de São Paulo: prefeito, vice, secretários municipais e subprefeitos, mês a mês.

Fonte: Portal de Dados Abertos da Prefeitura, "Histórico de Remuneração dos Servidores Ativos da Prefeitura de
São Paulo" (SIGPEC), um arquivo CSV por mês, com o nome de cada servidor da administração direta:
https://dados.prefeitura.sp.gov.br/dataset/remuneracao-servidores-prefeitura-de-sao-paulo

Colunas usadas: nome, cargo base, cargo em comissão, "Remuneração do Mês" (salário, verba de representação e o
que entra no teto), "Demais Elementos da Remuneração" (13º, férias, auxílio-refeição, atrasados e o que fica fora
do teto) e "Remuneração Bruta" (a soma). Quem tem decisão judicial para não aparecer não aparece (regra da
Prefeitura). A Prefeitura não publica gastos por pessoa (carro oficial, viagens): entra só o que cada um recebe.

Cada arquivo tem ~21 MB; o robô só baixa os meses que ainda não processou (ou que a Prefeitura publicou de novo)
e guarda só as linhas do prefeito, do vice, dos secretários e dos subprefeitos em dados/municipios/sp/.
Nomes de urna e partido do prefeito e do vice: TSE (2024). Fotos: Wikimedia Commons, só com licença livre.
"""
import csv
import io
import re
import time
from collections import Counter, defaultdict
from datetime import datetime

import pandas as pd

from .config import CACHE, DADOS, RAIZ
from .util import TempoEsgotado, _sessao, baixar, log, normalizar_nome, verificar_prazo

COD_IBGE = 3550308
INICIO = 202501  # mandato 2025–2028
PACOTE = "https://dados.prefeitura.sp.gov.br/api/3/action/package_show"
PACOTE_ID = "remuneracao-servidores-prefeitura-de-sao-paulo"
PAGINA = "https://dados.prefeitura.sp.gov.br/dataset/remuneracao-servidores-prefeitura-de-sao-paulo"
PASTA = DADOS / "municipios" / "sp"
LINHAS = PASTA / "prefeitura_remuneracao.csv"
ARQUIVOS = PASTA / "prefeitura_arquivos.csv"  # qual arquivo do Portal foi usado em cada mês
SAIDA = RAIZ / "site" / "dados" / "prefeituras.json"
FOTOS = RAIZ / "site" / "fotos"
TSE_ZIP = CACHE / "municipios" / "consulta_cand_2024.zip"
CARGOS = {"PREFEITO": "pr", "VICE PREFEITO": "vp", "SECRETARIO MUNICIPAL": "se", "SECRETARIO ESPECIAL": "se", "SUBPREFEITO": "sb"}
COLUNAS = ["aaaamm", "tp", "nome", "cargo_base", "cargo_comissao", "unidade", "pasta", "remuneracao_mes", "demais", "bruta", "excecao"]


# ---------------------------------------------------------------- arquivos do Portal
def _arquivos():
    """{AAAAMM: {id, url, quando}}: o arquivo mais recente de cada mês (às vezes a Prefeitura publica de novo).
    Prefere o CSV; se o mês só tem planilha (xlsx), usa a planilha."""
    itens = baixar(PACOTE, params={"id": PACOTE_ID}, timeout=120).json()["result"]["resources"]
    por_mes = {}
    for x in itens:
        url = x.get("url") or ""
        m = re.search(r"folha_(\d{6})", url)
        ext = url.lower().rsplit(".", 1)[-1]
        if not m or ext not in ("csv", "xlsx") or int(m.group(1)) < INICIO:
            continue
        am = int(m.group(1))
        chave = (ext == "csv", x.get("last_modified") or x.get("created") or "")
        if am not in por_mes or chave > por_mes[am]["chave"]:
            por_mes[am] = {"aaaamm": am, "id": x["id"], "url": url, "ext": ext, "chave": chave}
    return por_mes


def _num(t):
    """'34.700,46' ou '7,924.88' (um mês veio no formato americano) ou já um número (planilha)."""
    if isinstance(t, (int, float)):
        return round(float(t), 2)
    t = (t or "").strip()
    m = re.search(r"[.,](\d{1,2})$", t)
    inteiro = re.sub(r"\D", "", t[:m.start()] if m else t)
    try:
        return round(float(f"{inteiro or 0}.{m.group(1) if m else 0}"), 2)
    except ValueError:
        return 0.0


CACHE_BRUTO = None  # para testes: uma pasta onde guardar os CSVs inteiros (no robô, não guarda: são ~21 MB por mês)


def _decodificar(conteudo):
    """Os arquivos vêm em UTF-8, mas alguns meses vieram na página de código do DOS (cp850)."""
    try:
        return conteudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    for cod in ("cp850", "latin1"):
        texto = conteudo.decode(cod)
        if "REMUNERACAO" in normalizar_nome(texto[:400]):
            return texto
    return conteudo.decode("latin1")


def _ler_mes(am, url):
    """Baixa o CSV do mês e devolve só as linhas do prefeito, vice, secretários e subprefeitos."""
    verificar_prazo()
    guardado = CACHE_BRUTO / f"{am}.{url.lower().rsplit('.', 1)[-1]}" if CACHE_BRUTO else None
    if guardado and guardado.exists():
        conteudo = guardado.read_bytes()
    else:
        r = _sessao().get(url, timeout=600)
        r.raise_for_status()
        conteudo = r.content
        if guardado:
            guardado.parent.mkdir(parents=True, exist_ok=True)
            guardado.write_bytes(conteudo)
    if url.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        leitor = ([("" if v is None else v) for v in linha] for linha in wb.worksheets[0].iter_rows(values_only=True))
    else:
        texto = _decodificar(conteudo)
        primeira = texto[:texto.find("\n")]
        leitor = csv.reader(io.StringIO(texto), delimiter=";" if primeira.count(";") >= primeira.count(",") else ",")
    cab = [normalizar_nome(str(c)) for c in next(leitor)]
    col = lambda *partes: next(i for i, c in enumerate(cab) if all(p in c for p in partes))
    i_nome, i_base, i_com = col("NOME"), col("CARGO BASE"), col("CARGO EM COMISS")
    i_mes, i_dem, i_bruta, i_uni = col("REMUNERACAO DO M"), col("DEMAIS"), col("BRUTA"), col("UNIDADE")
    i_log, i_num, i_comp = col("LOG"), col("NUMERO"), col("COMPLEMENTO")
    i_exc = next((i for i, c in enumerate(cab) if "EXCE" in c), None)
    # endereço -> secretaria: para quem aparece só como "GABINETE DO SECRETARIO"
    por_endereco = defaultdict(Counter)
    achados = []
    for l in leitor:
        if len(l) <= max(i_uni, i_num):
            continue
        l = [v if isinstance(v, (int, float)) else str(v) for v in l]
        uni = str(l[i_uni]).strip()
        end = (normalizar_nome(str(l[i_log])), str(l[i_num]).strip())
        andar = (*end, normalizar_nome(str(l[i_comp])))
        if uni.startswith("SECRETARIA"):
            por_endereco[end][uni] += 1
            por_endereco[andar][uni] += 1
        base, com = str(l[i_base]).strip(), str(l[i_com]).strip()
        tp = CARGOS.get(com) or CARGOS.get(base)
        if tp:
            achados.append({"aaaamm": am, "tp": tp, "nome": str(l[i_nome]).strip(), "cargo_base": base, "cargo_comissao": com, "unidade": uni,
                            "_end": andar, "remuneracao_mes": _num(l[i_mes]), "demais": _num(l[i_dem]), "bruta": _num(l[i_bruta]),
                            "excecao": (str(l[i_exc]).strip() if i_exc is not None else "")})
    for a in achados:
        # palpite pelo endereço (mesmo andar; senão, mesmo prédio). O site só usa se não houver nada melhor.
        mais = por_endereco.get(a["_end"]) or por_endereco.get(a["_end"][:2])
        a["pasta"] = a["unidade"] if a["unidade"].startswith("SECRETARIA") or not mais else mais.most_common(1)[0][0]
        del a["_end"]
    return achados


def coletar_remuneracao():
    arquivos = _arquivos()
    feitos = pd.read_csv(ARQUIVOS).set_index("aaaamm")["id"].to_dict() if ARQUIVOS.exists() else {}
    linhas = pd.read_csv(LINHAS, dtype={"excecao": str}).fillna("") if LINHAS.exists() else pd.DataFrame(columns=COLUNAS)
    novos = [a for am, a in sorted(arquivos.items()) if feitos.get(am) != a["id"]]
    PASTA.mkdir(parents=True, exist_ok=True)
    try:
        for a in novos:
            log(f"Prefeitura SP: remuneração de {a['aaaamm'] % 100:02d}/{a['aaaamm'] // 100}")
            achados = _ler_mes(a["aaaamm"], a["url"])
            linhas = pd.concat([linhas[linhas.aaaamm != a["aaaamm"]], pd.DataFrame(achados, columns=COLUNAS)], ignore_index=True)
            feitos[a["aaaamm"]] = a["id"]
            linhas.sort_values(["aaaamm", "tp", "nome"]).to_csv(LINHAS, index=False)
            pd.DataFrame(sorted(feitos.items()), columns=["aaaamm", "id"]).to_csv(ARQUIVOS, index=False)
            time.sleep(1)
    except TempoEsgotado:
        raise
    return linhas


# ---------------------------------------------------------------- nomes com acento
_ACENTOS = {
    "EDUCACAO": "Educação", "SAUDE": "Saúde", "HABITACAO": "Habitação", "GESTAO": "Gestão", "JUSTICA": "Justiça",
    "SEGURANCA": "Segurança", "RELACOES": "Relações", "ASSIST": "Assistência", "DESENV": "Desenvolvimento",
    "ECONOMICO": "Econômico", "EFICIENCIA": "Eficiência", "DEFICIENCIA": "Deficiência", "CIDADAN": "Cidadania",
    "CIDADANIA": "Cidadania", "TRANSP": "Transporte", "TRANSPORTE": "Transporte", "OB": "Obras", "CRIATIV": "Criativa",
    "COMUNICACAO": "Comunicação", "INOVACAO": "Inovação", "ESTRATEGICOS": "Estratégicos", "POLITICAS": "Políticas",
    "PUBLICAS": "Públicas", "SE": "Sé", "SAO": "São", "BUTANTA": "Butantã", "JACANA": "Jaçanã", "TREMEMBE": "Tremembé",
    "BRASILANDIA": "Brasilândia", "O": "Ó", "LIMAO": "Limão", "JARAGUA": "Jaraguá", "CARRAO": "Carrão",
    "M": "M'", "ITAIM": "Itaim", "GUAIANASES": "Guaianases", "LICENCIAMENTO": "Licenciamento", "ADMINISTRACAO": "Administração",
    "GOVERNO": "Governo", "MUNICIPAL": "Municipal", "ESPECIAL": "Especial", "URBANA": "Urbana", "ECONOMIA": "Economia",
    "INTERNACIONAIS": "Internacionais", "INSTITUCIONAIS": "Institucionais", "ESPORTES": "Esportes", "LAZER": "Lazer",
    "MULHERES": "Mulheres", "PESSOA": "Pessoa", "ALIMENTAR": "Alimentar", "IGUALDADE": "Igualdade", "RACIAL": "Racial",
    "TECNOLOGIA": "Tecnologia", "CIENCIA": "Ciência", "FAZENDA": "Fazenda", "CASA": "Casa", "CIVIL": "Civil",
}
_PEQUENAS = {"DE", "DA", "DO", "DAS", "DOS", "E"}


def _bonito(texto):
    saida = []
    for i, w in enumerate(re.split(r"(\s+|/)", texto.strip())):
        if not w.strip() or w == "/":
            saida.append(w)
            continue
        u = w.upper()
        if u in _PEQUENAS and i:
            saida.append(u.lower() if u != "O" else "Ó")
        elif u in _ACENTOS:
            saida.append(_ACENTOS[u])
        else:
            saida.append(w.capitalize())
    t = re.sub(r"\bM'?\s?[Bb]oi\b", "M'Boi", "".join(saida))
    return re.sub(r"\bDo Ó\b", "do Ó", t)


# Secretários que aparecem na folha só como "GABINETE DO SECRETARIO" (sem o nome da secretaria).
# Conferido no site da Prefeitura; quem aparecer novo assim cai no palpite pelo endereço e o robô avisa.
_PASTAS_CONHECIDAS = {
    "ELISABETE FRANCA": "SECRETARIA MUNICIPAL DE URBANISMO E LICENCIAMENTO",
    "LUIZ CARLOS ZAMARCO": "SECRETARIA MUNICIPAL DA SAUDE",
}
_UNIDADES = {"CASA CIVIL": "SECRETARIA MUNICIPAL DA CASA CIVIL", "GABINETE DO PREFEITO": "GABINETE DO PREFEITO"}


def _resolver_pastas(g):
    """Secretaria de cada mês: a da folha; senão, a da tabela acima; senão, a da própria pessoa em outro mês;
    senão, o palpite pelo endereço."""
    nome = normalizar_nome(g.nome.iloc[0])
    certas = [u for u in g.unidade if u.startswith("SECRETARIA")]
    saida = []
    for r in g.itertuples():
        if r.tp != "se" or r.unidade.startswith("SECRETARIA"):
            saida.append(r.unidade)
        elif r.unidade in _UNIDADES:
            saida.append(_UNIDADES[r.unidade])
        elif nome in _PASTAS_CONHECIDAS:
            saida.append(_PASTAS_CONHECIDAS[nome])
        elif certas:
            saida.append(Counter(certas).most_common(1)[0][0])
        else:
            log(f"  Prefeitura SP: secretaria de {r.nome} em {r.aaaamm} pelo endereço ({r.pasta}); confira e acrescente em _PASTAS_CONHECIDAS")
            saida.append(r.pasta)
    return saida


def _pasta(tp, unidade):
    """'SECRETARIA MUNICIPAL DE EDUCACAO' -> 'Secretaria Municipal de Educação'; 'SUBPREFEITURA LAPA' -> 'Lapa'."""
    if tp in ("pr", "vp"):
        return "Prefeitura de São Paulo"
    if tp == "sb":
        return _bonito(re.sub(r"^SUBPREFEITURA\s+", "", unidade))
    return _bonito(unidade) if unidade else "Secretaria municipal"


# nomes femininos que não terminam em "a" e masculinos que terminam (ajustar quando aparecer alguém novo)
_FEMININOS = {"ELISABETE", "ERIKA", "CIBELE", "REGINA", "SILVIA", "JULIANA", "ANGELA", "ELIANA", "MARCELA", "LUCIANA", "ANA", "RAQUEL",
              "BEATRIZ", "ISABEL", "IRIS", "LIZ", "INES", "ALICE", "CRISTIANE", "DENISE", "ELAINE", "GISELE", "LILIAN", "MIRIAM", "RUTH",
              "SUELI", "SIMONE", "SOLANGE", "VIVIANE", "JAQUELINE", "ALINE", "ROSE", "ROSELI", "ELIZABETH", "CARMEN", "KARIN",
              "EUNICE", "ELISETE", "DAMARIS", "THAMYRIS", "IRENE", "JANETE", "CLARICE", "LUCIENE", "MICHELE", "NOEMI"}
_MASCULINOS = {"LUCA", "JOSHUA", "NICOLA", "BATISTA", "GARCIA"}


def _feminino(nome):
    p = normalizar_nome(nome).split()[0] if nome else ""
    if p in _FEMININOS:
        return True
    if p in _MASCULINOS:
        return False
    return p.endswith("A")


def _cargo_nome(tp, pasta, fem):
    if tp == "pr":
        return "Prefeita de São Paulo" if fem else "Prefeito de São Paulo"
    if tp == "vp":
        return "Vice-prefeita de São Paulo" if fem else "Vice-prefeito de São Paulo"
    if tp == "sb":
        return f"{'Subprefeita' if fem else 'Subprefeito'} ({pasta})"
    if not pasta.startswith("Secretaria"):  # secretário especial lotado, por exemplo, no Gabinete do Prefeito
        return f"{'Secretária' if fem else 'Secretário'} Especial ({pasta})"
    return re.sub(r"^Secretaria\b", "Secretária" if fem else "Secretário", pasta)


# ---------------------------------------------------------------- TSE: prefeito e vice
def _eleitos_executivo():
    """{nome civil normalizado: (nome de urna, partido)} do prefeito e do vice eleitos em 2024 em São Paulo."""
    import zipfile
    saida = {}
    if not TSE_ZIP.exists():
        return saida
    with zipfile.ZipFile(TSE_ZIP) as z, z.open("consulta_cand_2024_SP.csv") as f:
        for l in csv.DictReader(io.TextIOWrapper(f, encoding="latin1"), delimiter=";"):
            if l["NM_UE"] == "SÃO PAULO" and l["CD_CARGO"] in ("11", "12") and l["DS_SIT_TOT_TURNO"].startswith("ELEITO"):
                saida[normalizar_nome(l["NM_CANDIDATO"])] = (l["NM_URNA_CANDIDATO"].strip(), l["SG_PARTIDO"])
            elif l["NM_UE"] == "SÃO PAULO" and l["CD_CARGO"] == "12" and l["SG_PARTIDO"] and l["DS_SITUACAO_CANDIDATURA"] == "APTO":
                # o vice não tem "eleito" no turno: guarda todos os vices aptos; só o que estiver na folha é usado
                saida.setdefault(normalizar_nome(l["NM_CANDIDATO"]), (l["NM_URNA_CANDIDATO"].strip(), l["SG_PARTIDO"]))
    return saida


_MINUSCULAS = {"da", "de", "do", "das", "dos", "e"}
# a folha vem sem acentos; estes nomes quase sempre têm (os que às vezes não têm, como Luis e Angela, ficam como estão)
_NOMES_ACENTO = {
    "jose": "José", "joao": "João", "andre": "André", "antonio": "Antônio", "fabricio": "Fabrício", "sergio": "Sérgio",
    "rogerio": "Rogério", "marcio": "Márcio", "flavia": "Flávia", "katia": "Kátia", "decio": "Décio", "alvaro": "Álvaro",
    "vinicius": "Vinícius", "valerio": "Valério", "conceicao": "Conceição", "falcao": "Falcão", "antao": "Antão",
    "franca": "França", "celia": "Célia", "araujo": "Araújo", "goncalves": "Gonçalves", "simoes": "Simões",
    "patricia": "Patrícia", "lucia": "Lúcia", "julio": "Júlio", "claudio": "Cláudio", "claudia": "Cláudia",
    "fabio": "Fábio", "flavio": "Flávio", "mario": "Mário", "otavio": "Otávio", "helio": "Hélio", "barbara": "Bárbara", "sonia": "Sônia",
    "vania": "Vânia", "monica": "Mônica", "cicero": "Cícero", "sebastiao": "Sebastião", "estevao": "Estevão",
    "assuncao": "Assunção", "romao": "Romão", "simao": "Simão", "galvao": "Galvão", "brandao": "Brandão", "junior": "Júnior",
}


def _titulo(nome):
    return " ".join(w if w in _MINUSCULAS and i else _NOMES_ACENTO.get(w, w.capitalize()) for i, w in enumerate(nome.lower().split()))


# ---------------------------------------------------------------- fotos (Wikimedia Commons, licença livre)
def _fotos(pessoas):
    """Só para prefeito, vice e secretários; usa a mesma busca e as mesmas regras do governo federal."""
    from . import fotos as F
    antigo = F.CARGO_OK
    F.CARGO_OK = re.compile(r"prefeit|mayor|secret", re.I)
    try:
        return F._governo_commons([{"id": p["id"], "nome": p["n"], "nome_civil": p["nc"], "casa": "executivo"}
                                   for p in pessoas if p["tp"] in ("pr", "vp", "se")])
    finally:
        F.CARGO_OK = antigo


# ---------------------------------------------------------------- site/dados/prefeituras.json
def _r(v):
    return int(round(float(v)))


def site(linhas, fotos=True):
    if not len(linhas):
        return None
    linhas = linhas.copy()
    linhas["aaaamm"] = linhas.aaaamm.astype(int)
    ultimo = int(linhas.aaaamm.max())
    anos = sorted({str(a // 100) for a in linhas.aaaamm})
    tse = _eleitos_executivo()
    # vereador de SP com o mesmo nome completo: liga as duas páginas
    vereadores = {}
    arq_ver = PASTA / "vereadores.csv"
    if arq_ver.exists():
        v = pd.read_csv(arq_ver).fillna("")
        vereadores = {normalizar_nome(n): (int(c), nome) for c, n, nome in zip(v.codigo, v.nome_civil, v.nome) if n}
    pessoas = []
    # mesma pessoa escrita com e sem "da"/"de" em meses diferentes: junta pelo nome sem essas palavras
    chave = lambda n: " ".join(w for w in normalizar_nome(n).split() if w not in _PEQUENAS)
    for _, g in linhas.assign(chave=linhas.nome.map(chave)).groupby("chave"):
        nome_n = normalizar_nome(g.sort_values("aaaamm").nome.iloc[-1])
        g = g.sort_values("aaaamm").copy()
        g["pasta"] = _resolver_pastas(g)
        ult = g.iloc[-1]
        tp = ult.tp
        fem = _feminino(ult.nome)
        pasta = _pasta(tp, ult.pasta)
        urna, partido = tse.get(nome_n, (None, None))
        ver = vereadores.get(nome_n)
        exibido = _titulo(urna) if urna else ver[1] if ver else _titulo(ult.nome)
        mes = g.groupby("aaaamm")[["remuneracao_mes", "demais", "bruta"]].sum()
        # quem saiu: os acertos do mês da saída (férias, 13º proporcional...) ficam fora das médias, à parte
        # (só o que passa do normal dele, como o auxílio-refeição, e se passar de R$ 3 mil)
        saida = 0.0
        if int(ult.aaaamm) != ultimo:
            normal = float(mes.demais.iloc[:-1].median()) if len(mes) > 1 else 0.0
            excesso = float(mes.demais.iloc[-1]) - normal
            if excesso >= 3000:
                saida = excesso
                mes.loc[mes.index[-1], ["demais", "bruta"]] = [normal, mes.remuneracao_mes.iloc[-1] + normal]
        serie = [[int(am), _r(r.bruta), 0, 0, 0, 0] for am, r in mes.iterrows()]

        def bloco(filtro):
            s = mes[[filtro(am) for am in mes.index]]
            if not len(s):
                return None
            m = len(s)
            cats = {k: _r(v) for k, v in (("salario", s.remuneracao_mes.sum()), ("outros_rendimentos", s.demais.sum())) if _r(v)}
            return {"m": m, "mg": int((s.bruta > 0).sum()), "mc": 0, "me": 0, "g": _r(s.bruta.sum()), "c": 0, "e": 0,
                    "pm": 0, "mp": 0, "pu": 0, "ep": 0, "cats": cats}

        per = {}
        for a in anos:
            b = bloco(lambda am, a=a: str(am // 100) == a)
            if b:
                per[a] = b
        per["leg"] = bloco(lambda am: True)
        # cargos ocupados (em ordem), com o primeiro e o último mês de cada um
        cargos = []
        for am, gm in g.groupby("aaaamm"):
            nomes = list(dict.fromkeys(_cargo_nome(r.tp, _pasta(r.tp, r.pasta), fem) for r in gm.itertuples()))
            if cargos and cargos[-1][0] in nomes:  # continua no mesmo cargo: ele primeiro
                nomes.remove(cargos[-1][0])
                cargos[-1][2] = int(am)
            for nome_c in nomes:
                cargos.append([nome_c, int(am), int(am)])
        pid = f"pre-{COD_IBGE}-{re.sub(r'[^a-z0-9]+', '-', nome_n.lower()).strip('-')}"
        # servidor cedido por outro órgão, que paga o salário (exceções 2 e 3 da Prefeitura): a folha mostra só uma parte
        exc = g.excecao.astype(str).str.strip()
        cedido = int(exc.isin(["2", "3"]).sum())
        foto = FOTOS / f"{pid}.webp"
        pessoas.append({
            "id": pid, "k": "p", "cid": COD_IBGE, "tp": tp, "n": exibido, "nc": _titulo(ult.nome),
            "g": _cargo_nome(tp, pasta, fem), "pa": pasta, "pt": partido, "uf": "SP",
            "f": f"fotos/{pid}.webp" if foto.exists() else None, "x": 1 if int(ult.aaaamm) == ultimo else 0,
            "o": PAGINA, "per": per, "t": serie, "dt": {}, "cg": cargos,
            **({"rel": f"ver-{COD_IBGE}-{ver[0]}"} if ver else {}),
            **({"ced": cedido} if cedido else {}),
            **({"q": [1, _r(saida)]} if saida >= 1 else {}),
        })
    if fotos:
        try:
            novas = _fotos([p for p in pessoas if not p.get("rel")])
            if novas:
                log(f"  {novas} fotos novas da Prefeitura (Wikimedia Commons)")
        except TempoEsgotado:
            raise
        except Exception as e:  # noqa: BLE001 — foto é opcional
            log(f"  Fotos da Prefeitura: {e}")
    creditos = {}
    from .fotos import CREDITOS
    if CREDITOS.exists():
        import json as _json
        creditos = _json.loads(CREDITOS.read_text(encoding="utf-8")).get("fotos", {})
    for p in pessoas:
        if (FOTOS / f"{p['id']}.webp").exists():
            p["f"] = f"fotos/{p['id']}.webp"
            c = creditos.get(p["id"])
            if c:
                p["fc"] = {"a": c.get("autor"), "l": c.get("licenca"), "u": c.get("pagina")}
        elif p.get("rel") and (FOTOS / f"{p['rel']}.webp").exists():
            # quem também é vereador usa a foto oficial da Câmara Municipal
            p["f"] = f"fotos/{p['rel']}.webp"
            p["fc"] = {"a": "Câmara Municipal de São Paulo", "u": "https://www.saopaulo.sp.leg.br/vereadores/membros/"}
    ordem = {"pr": 0, "vp": 1, "se": 2, "sb": 3}
    pessoas.sort(key=lambda p: (ordem[p["tp"]], normalizar_nome(p["n"])))
    dados = {
        "meta": {
            "gerado_em": datetime.now().isoformat(timespec="seconds"), "tipos": [], "categorias": {},
            "cidades": {str(COD_IBGE): {"n": "São Paulo", "uf": "SP", "casa": "Prefeitura de São Paulo", "inicio": INICIO,
                                        "ultimo_mes": ultimo, "anos": anos, "fonte": PAGINA}},
        },
        "p": pessoas,
    }
    import json
    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    SAIDA.write_text(json.dumps(dados, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    n = Counter(p["tp"] for p in pessoas if p["x"])
    log(f"Site: {SAIDA.relative_to(RAIZ)} ({SAIDA.stat().st_size / 1e3:.0f} KB; no cargo em {ultimo % 100:02d}/{ultimo // 100}: "
        f"{n['pr']} prefeito, {n['vp']} vice, {n['se']} secretários, {n['sb']} subprefeitos)")
    return dados


def coletar():
    linhas = coletar_remuneracao()
    site(linhas)


def executar_site():
    """Só remonta site/dados/prefeituras.json com o que já está em dados/municipios/sp/ (sem baixar nada)."""
    if LINHAS.exists():
        site(pd.read_csv(LINHAS, dtype={"excecao": str}).fillna(""), fotos=False)
