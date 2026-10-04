"""Prefeitura de São Paulo: prefeito, vice, secretários municipais e subprefeitos, mês a mês.

Fonte: Portal de Dados Abertos da Prefeitura, "Histórico de Remuneração dos Servidores Ativos da Prefeitura de
São Paulo" (SIGPEC), um arquivo CSV por mês, com o nome de cada servidor da administração direta:
https://dados.prefeitura.sp.gov.br/dataset/remuneracao-servidores-prefeitura-de-sao-paulo

Colunas usadas: nome, cargo base, cargo em comissão, "Remuneração do Mês" (salário, verba de representação e o
que entra no teto), "Demais Elementos da Remuneração" (13º, férias, auxílio-refeição, atrasados e o que fica fora
do teto) e "Remuneração Bruta" (a soma). Quem tem decisão judicial para não aparecer não aparece (regra da
Prefeitura). A Prefeitura não publica gastos por pessoa (carro oficial, viagens): entra só o que cada um recebe.

Cada arquivo tem ~21 MB; o robô só baixa os meses que ainda não processou (ou que a Prefeitura publicou de novo)
e guarda só as linhas do prefeito, do vice, dos secretários e dos subprefeitos em dados/municipios/sp/.
Nomes de urna e partido do prefeito e do vice: TSE (2024). Fotos: Wikimedia Commons, só com licença livre.
"""
import csv
import io
import re
import time
from collections import Counter, defaultdict

import pandas as pd

from ..config import DADOS
from ..util import TempoEsgotado, _sessao, baixar, gravar_csv, log, normalizar_nome, verificar_prazo
from . import comum

COD_IBGE = 3550308
INICIO = 202501  # mandato 2025–2028
PACOTE = "https://dados.prefeitura.sp.gov.br/api/3/action/package_show"
PACOTE_ID = "remuneracao-servidores-prefeitura-de-sao-paulo"
PAGINA = "https://dados.prefeitura.sp.gov.br/dataset/remuneracao-servidores-prefeitura-de-sao-paulo"
PASTA = DADOS / "municipios" / "sp"
LINHAS = PASTA / "prefeitura_remuneracao.csv"
ARQUIVOS = PASTA / "prefeitura_arquivos.csv"  # qual arquivo do Portal foi usado em cada mês
CARGOS = {"PREFEITO": "pr", "VICE PREFEITO": "vp", "SECRETARIO MUNICIPAL": "se", "SECRETARIO ESPECIAL": "se", "SUBPREFEITO": "sb"}
COLUNAS = ["aaaamm", "tp", "nome", "cargo_base", "cargo_comissao", "unidade", "pasta", "remuneracao_mes", "demais", "bruta", "excecao"]


# ---------------------------------------------------------------- arquivos do Portal
def _arquivos():
    """{AAAAMM: {id, url, quando}}: o arquivo mais recente de cada mês (às vezes a Prefeitura publica de novo).
    Prefere o CSV; se o mês só tem planilha (xlsx), usa a planilha."""
    itens = baixar(PACOTE, params={"id": PACOTE_ID}, timeout=120).json()["result"]["resources"]
    por_mes = {}
    for x in itens:
        url = x.get("url") or ""
        m = re.search(r"folha_(\d{6})", url)
        ext = url.lower().rsplit(".", 1)[-1]
        if not m or ext not in ("csv", "xlsx") or int(m.group(1)) < INICIO:
            continue
        am = int(m.group(1))
        chave = (ext == "csv", x.get("last_modified") or x.get("created") or "")
        if am not in por_mes or chave > por_mes[am]["chave"]:
            por_mes[am] = {"aaaamm": am, "id": x["id"], "url": url, "ext": ext, "chave": chave}
    return por_mes


_num = comum.num


CACHE_BRUTO = None  # para testes: uma pasta onde guardar os CSVs inteiros (no robô, não guarda: são ~21 MB por mês)


def _decodificar(conteudo):
    """Os arquivos vêm em UTF-8, mas alguns meses vieram na página de código do DOS (cp850)."""
    try:
        return conteudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    for cod in ("cp850", "latin1"):
        texto = conteudo.decode(cod)
        if "REMUNERACAO" in normalizar_nome(texto[:400]):
            return texto
    return conteudo.decode("latin1")


def _ler_mes(am, url):
    """Baixa o CSV do mês e devolve só as linhas do prefeito, vice, secretários e subprefeitos."""
    verificar_prazo()
    guardado = CACHE_BRUTO / f"{am}.{url.lower().rsplit('.', 1)[-1]}" if CACHE_BRUTO else None
    if guardado and guardado.exists():
        conteudo = guardado.read_bytes()
    else:
        r = _sessao().get(url, timeout=600)
        r.raise_for_status()
        conteudo = r.content
        if guardado:
            guardado.parent.mkdir(parents=True, exist_ok=True)
            guardado.write_bytes(conteudo)
    if url.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        leitor = ([("" if v is None else v) for v in linha] for linha in wb.worksheets[0].iter_rows(values_only=True))
    else:
        texto = _decodificar(conteudo)
        primeira = texto[:texto.find("\n")]
        leitor = csv.reader(io.StringIO(texto), delimiter=";" if primeira.count(";") >= primeira.count(",") else ",")
    cab = [normalizar_nome(str(c)) for c in next(leitor)]
    col = lambda *partes: next(i for i, c in enumerate(cab) if all(p in c for p in partes))
    i_nome, i_base, i_com = col("NOME"), col("CARGO BASE"), col("CARGO EM COMISS")
    i_mes, i_dem, i_bruta, i_uni = col("REMUNERACAO DO M"), col("DEMAIS"), col("BRUTA"), col("UNIDADE")
    i_log, i_num, i_comp = col("LOG"), col("NUMERO"), col("COMPLEMENTO")
    i_exc = next((i for i, c in enumerate(cab) if "EXCE" in c), None)
    # endereço -> secretaria: para quem aparece só como "GABINETE DO SECRETARIO"
    por_endereco = defaultdict(Counter)
    achados = []
    for l in leitor:
        if len(l) <= max(i_uni, i_num):
            continue
        l = [v if isinstance(v, (int, float)) else str(v) for v in l]
        uni = str(l[i_uni]).strip()
        end = (normalizar_nome(str(l[i_log])), str(l[i_num]).strip())
        andar = (*end, normalizar_nome(str(l[i_comp])))
        if uni.startswith("SECRETARIA"):
            por_endereco[end][uni] += 1
            por_endereco[andar][uni] += 1
        base, com = str(l[i_base]).strip(), str(l[i_com]).strip()
        tp = CARGOS.get(com) or CARGOS.get(base)
        if tp:
            achados.append({"aaaamm": am, "tp": tp, "nome": str(l[i_nome]).strip(), "cargo_base": base, "cargo_comissao": com, "unidade": uni,
                            "_end": andar, "remuneracao_mes": _num(l[i_mes]), "demais": _num(l[i_dem]), "bruta": _num(l[i_bruta]),
                            "excecao": (str(l[i_exc]).strip() if i_exc is not None else "")})
    for a in achados:
        # palpite pelo endereço (mesmo andar; senão, mesmo prédio). O site só usa se não houver nada melhor.
        mais = por_endereco.get(a["_end"]) or por_endereco.get(a["_end"][:2])
        a["pasta"] = a["unidade"] if a["unidade"].startswith("SECRETARIA") or not mais else mais.most_common(1)[0][0]
        del a["_end"]
    return achados


def coletar_remuneracao():
    arquivos = _arquivos()
    feitos = pd.read_csv(ARQUIVOS).set_index("aaaamm")["id"].to_dict() if ARQUIVOS.exists() else {}
    linhas = pd.read_csv(LINHAS, dtype={"excecao": str}).fillna("") if LINHAS.exists() else pd.DataFrame(columns=COLUNAS)
    novos = [a for am, a in sorted(arquivos.items()) if feitos.get(am) != a["id"]]
    PASTA.mkdir(parents=True, exist_ok=True)
    try:
        for a in novos:
            log(f"Prefeitura SP: remuneração de {a['aaaamm'] % 100:02d}/{a['aaaamm'] // 100}")
            achados = _ler_mes(a["aaaamm"], a["url"])
            linhas = pd.concat([linhas[linhas.aaaamm != a["aaaamm"]], pd.DataFrame(achados, columns=COLUNAS)], ignore_index=True)
            if not gravar_csv(linhas.sort_values(["aaaamm", "tp", "nome"]), LINHAS):
                break  # recusado por perda de cobertura (util.gravar_com): fica o que estava, e o mês é lido de novo
            feitos[a["aaaamm"]] = a["id"]
            gravar_csv(pd.DataFrame(sorted(feitos.items()), columns=["aaaamm", "id"]), ARQUIVOS)
            time.sleep(1)
    except TempoEsgotado:
        raise
    return linhas


# Secretários que aparecem na folha só como "GABINETE DO SECRETARIO" (sem o nome da secretaria).
# Conferido no site da Prefeitura; quem aparecer novo assim cai no palpite pelo endereço e o robô avisa.
_PASTAS_CONHECIDAS = {
    "ELISABETE FRANCA": "SECRETARIA MUNICIPAL DE URBANISMO E LICENCIAMENTO",
    "LUIZ CARLOS ZAMARCO": "SECRETARIA MUNICIPAL DA SAUDE",
}
_UNIDADES = {"CASA CIVIL": "SECRETARIA MUNICIPAL DA CASA CIVIL", "GABINETE DO PREFEITO": "GABINETE DO PREFEITO"}


def _resolver_pastas(g):
    """Secretaria de cada mês: a da folha; senão, a da tabela acima; senão, a da própria pessoa em outro mês;
    senão, o palpite pelo endereço."""
    nome = normalizar_nome(g.nome.iloc[0])
    certas = [u for u in g.unidade if u.startswith("SECRETARIA")]
    saida = []
    for r in g.itertuples():
        if r.tp != "se" or r.unidade.startswith("SECRETARIA"):
            saida.append(r.unidade)
        elif r.unidade in _UNIDADES:
            saida.append(_UNIDADES[r.unidade])
        elif nome in _PASTAS_CONHECIDAS:
            saida.append(_PASTAS_CONHECIDAS[nome])
        elif certas:
            saida.append(Counter(certas).most_common(1)[0][0])
        else:
            log(f"  Prefeitura SP: secretaria de {r.nome} em {r.aaaamm} pelo endereço ({r.pasta}); confira e acrescente em _PASTAS_CONHECIDAS")
            saida.append(r.pasta)
    return saida


def _pasta(tp, unidade):
    """'SECRETARIA MUNICIPAL DE EDUCACAO' -> 'Secretaria Municipal de Educação'; 'SUBPREFEITURA LAPA' -> 'Lapa'."""
    if tp in ("pr", "vp"):
        return "Prefeitura de São Paulo"
    if tp == "sb":
        return comum.bonito(re.sub(r"^SUBPREFEITURA\s+", "", unidade))
    return comum.bonito(unidade) if unidade else "Secretaria municipal"



def _vereadores_sp():
    """Todos os vereadores de SP desde 2025, inclusive quem se licenciou antes de ter página (vai para o secretariado)."""
    arq = PASTA / "vereadores.csv"
    if not arq.exists():
        return {}
    v = pd.read_csv(arq).fillna("")
    return {normalizar_nome(n): (f"ver-{COD_IBGE}-{int(c)}", nome) for c, n, nome in zip(v.codigo, v.nome_civil, v.nome) if n}


CFG = {
    "especial": True,  # secretário fora de uma secretaria (no Gabinete do Prefeito) é "Secretário Especial"
    "cod": COD_IBGE, "n": "São Paulo", "uf": "SP", "de": "de São Paulo", "casa": "Prefeitura de São Paulo", "inicio": INICIO,
    "fonte": PAGINA, "pagina": PAGINA,
    "salario_nota": ("Remuneração bruta da folha da Prefeitura: salário e verba de representação, mais 13º, férias, auxílio-refeição e "
                     "pagamentos atrasados (que a Prefeitura chama de \"demais elementos\"). Valores antes do desconto de impostos."),
    "notas": ["A folha é de todos os servidores da administração direta, com o nome, publicada todo mês nos dados abertos da Prefeitura.",
              "Quem tem decisão judicial para não aparecer na folha não aparece aqui."],
    "credito_camara": "Câmara Municipal de São Paulo", "pagina_camara": "https://www.saopaulo.sp.leg.br/vereadores/membros/",
    "vereadores": _vereadores_sp,
}


def montar():
    """As linhas de dados/municipios/sp/prefeitura_remuneracao.csv no formato comum (secretaria resolvida por pessoa)."""
    if not LINHAS.exists():
        return None
    linhas = pd.read_csv(LINHAS, dtype={"excecao": str}).fillna("")
    if not len(linhas):
        return None
    saida = []
    chave = lambda n: " ".join(w for w in normalizar_nome(n).split() if w not in comum._PEQUENAS)
    for _, g in linhas.assign(chave=linhas.nome.map(chave)).groupby("chave"):
        g = g.sort_values("aaaamm").copy()
        g["pasta"] = _resolver_pastas(g)
        for r in g.itertuples():
            saida.append({"aaaamm": int(r.aaaamm), "tp": r.tp, "nome": r.nome, "pasta": _pasta(r.tp, r.pasta),
                          "salario": float(r.remuneracao_mes), "decimo": 0.0, "outros": float(r.demais), "bruta": float(r.bruta),
                          # servidor cedido por outro órgão, que paga o salário (exceções 2 e 3 da Prefeitura): a folha mostra só uma parte
                          "cedido": 1 if str(r.excecao).strip() in ("2", "3") else 0})
    return comum.montar(CFG, pd.DataFrame(saida, columns=comum.COLUNAS))


# O robots.txt do Portal de Dados Abertos da Prefeitura (dados.prefeitura.sp.gov.br) tem "Disallow: /" para todos os
# robôs. O portal está na lista de exceções (coleta/util.py, regra no CLAUDE.md): a folha é dado que a LAI manda abrir.
# Se a Prefeitura pedir para parar ou bloquear, troque para True: o site fica com o que já estava gravado.
BLOQUEADO_ROBOTS = False


def coletar():
    if BLOQUEADO_ROBOTS:
        log("  Prefeitura de São Paulo: o robots.txt do portal não permite robôs; fica o que já estava gravado")
        return
    coletar_remuneracao()
